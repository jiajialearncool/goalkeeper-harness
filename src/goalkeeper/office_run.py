"""Real F1 execution, separate from the deterministic office-format development fixture."""

from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

from .auth import AuthBroker
from .engine import Engine
from .evaluation import NotReady, read, office_attempt_count
from .files import digest, snapshot, EXCLUDED
from .grader import save_report
from .models import CheckSpec, Goal, RunConfig, RunState, uid
from .office import RUNTIME
from .office_finalize import finalize_office_delivery
from .paths import DATA, PROJECT
from .provider import ResponsesProvider
from .store import Store

NAMES = ["销售分析.xlsx", "销售汇报-六页初版.pptx", "销售汇报-三页改版.pptx"]
REVISION = "用户改版要求：保留销售分析.xlsx 和六页初版，另制作销售汇报-三页改版.pptx，恰好三页，依次说明净销售额趋势、退款趋势和产品构成。全部图表使用可编辑原生图表。核对表格和两份材料的数据一致，再运行最终检查。"


def source_rows():
    return [
        [f"{m}月", ["基础款", "标准款", "专业款"][p], 70 + m * 10 + p * 15, [80, 120, 200][p], (m + p) * 20]
        for m in range(1, 7)
        for p in range(3)
    ]


def inspect_outputs(root: Path, *, revision=True) -> dict:
    expected = source_rows()
    total = sum(r[2] * r[3] - r[4] for r in expected)
    checks = {}
    try:
        cached = load_workbook(root / NAMES[0], data_only=True)
        formula = load_workbook(root / NAMES[0], data_only=False)
        raw = cached["原始销售数据"]
        checks["source_rows_unchanged"] = [
            list(r) for r in raw.iter_rows(min_row=2, max_row=19, max_col=5, values_only=True)
        ] == expected
        checks["correct_total"] = cached["销售分析"]["D12"].value == total
        checks["editable_formulas"] = (
            all(
                formula["原始销售数据"][f"{column}{row}"].data_type == "f"
                for row in range(2, 20)
                for column in ["F", "G"]
            )
            and formula["销售分析"]["D12"].data_type == "f"
        )
        checks["native_workbook_chart"] = bool(formula["销售分析"]._charts)
        wanted = [(NAMES[1], 6)] + ([(NAMES[2], 3)] if revision else [])
        for name, count in wanted:
            with zipfile.ZipFile(root / name) as file:
                slides = [
                    n for n in file.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")
                ]
                charts = [n for n in file.namelist() if "/charts/chart" in n and n.endswith(".xml")]
                checks[name + "/slides"] = len(slides) == count
                checks[name + "/native_charts"] = len(charts) >= 2
                chart_values = {
                    float(e.text)
                    for n in charts
                    for e in ET.fromstring(file.read(n)).iter()
                    if e.tag.endswith("}v") and e.text and _numeric(e.text)
                }
                monthly = [sum(r[2] * r[3] - r[4] for r in expected if r[0] == f"{m}月") for m in range(1, 7)]
                checks[name + "/monthly_data"] = set(monthly).issubset(chart_values)
                mix = [
                    sum(r[2] * r[3] - r[4] for r in expected if r[1] == product)
                    for product in ["基础款", "标准款", "专业款"]
                ]
                checks[name + "/product_data"] = set(mix).issubset(chart_values)
                if name == NAMES[2]:
                    refunds = [sum(r[4] for r in expected if r[0] == f"{m}月") for m in range(1, 7)]
                    checks[name + "/refund_data"] = set(refunds).issubset(chart_values)
        cached.close()
        formula.close()
    except Exception as error:
        checks["file_integrity"] = False
        return {
            "passed": False,
            "checks": checks,
            "error": type(error).__name__ + ": " + str(error),
            "total": total,
        }
    return {"passed": all(checks.values()), "checks": checks, "total": total}


def _numeric(value):
    try:
        float(value)
        return True
    except ValueError:
        return False


def editable_input_check(root: Path) -> dict:
    """Import and recalculate in a trusted, independent process; never execute candidate scripts."""
    build = PROJECT / "work" / ("office-verify-" + uid())
    build.mkdir(parents=True)
    (build / "node_modules").symlink_to(RUNTIME / "node/node_modules", target_is_directory=True)
    program = """import {SpreadsheetFile,FileBlob} from '@oai/artifact-tool';
const wb=await SpreadsheetFile.importXlsx(await FileBlob.load(process.argv[2]));
const raw=wb.worksheets.getItem('原始销售数据'),summary=wb.worksheets.getItem('销售分析');
wb.recalculate();const before=summary.getRange('D12').values[0][0];const original=raw.getRange('C2').values[0][0];const price=raw.getRange('D2').values[0][0];raw.getRange('C2').values=[[original+1]];wb.recalculate();const after=summary.getRange('D12').values[0][0];console.log(JSON.stringify({passed:after===before+price,before,after,price}));if(after!==before+price)process.exit(1);
"""
    (build / "verify.mjs").write_text(program)
    run = subprocess.run(
        [str(RUNTIME / "node/bin/node"), str(build / "verify.mjs"), str(root / NAMES[0])],
        capture_output=True,
        text=True,
        timeout=120,
    )
    try:
        result = json.loads(run.stdout.strip().splitlines()[-1])
    except (ValueError, IndexError):
        result = {"passed": False, "error": (run.stderr + run.stdout)[-3000:]}
    return {**result, "exit_code": run.returncode}


async def office_pilot(diagnostic_reason: str | None = None) -> dict:
    broker = AuthBroker(DATA / "private")
    provider = ResponsesProvider(broker)
    store = Store(DATA / "private")
    manifest_path = DATA / "private/f1-run.json"
    try:
        if not broker.status()["connected"]:
            raise NotReady("请先授权 GoalKeeper，再运行真实 F1 演示。")
        if not read(PROJECT / "outputs/preflight.json").get("real_smoke_passed"):
            raise NotReady("先通过真实模型与工具接入检查，再运行 F1。")
        manifest = read(manifest_path)
        if diagnostic_reason:
            if not manifest or store.get(manifest["run_id"]).phase not in {"failed", "budget_exhausted"}:
                raise NotReady("只有保留的失败尝试才能进入 F1 诊断。")
            if len(read(DATA / "private/pilot-ledger.json", [])) + office_attempt_count() >= 24:
                raise NotReady("pilot 与 F1 的24次尝试预算已经用完。")
            previous = store.get(manifest["run_id"])
            saved = read(PROJECT / "outputs/f1-validation.json")
            save_report(PROJECT / f"outputs/f1-attempt-{previous.id}.json", saved)
            identifier = uid()
            root = DATA / "runs" / identifier / "workspace"
            shutil.copytree(previous.root, root, ignore=shutil.ignore_patterns(*EXCLUDED))
            (root / "node_modules").symlink_to(RUNTIME / "node/node_modules", target_is_directory=True)
            store.create(
                RunState(
                    id=identifier,
                    workspace=str(root),
                    goal=previous.goal.model_copy(deep=True),
                    config=RunConfig(),
                    evaluation_phase="demo",
                    final_checks=previous.final_checks,
                )
            )
            if store.has_event(previous.id, "f1_revision_requested"):
                store.event(
                    identifier,
                    "f1_revision_requested",
                    {
                        "source_run_id": previous.id,
                        "user_feedback": REVISION,
                        "inherited_actual_feedback": True,
                    },
                )
            manifest = {
                **manifest,
                "run_id": identifier,
                "prior_run_ids": [*manifest.get("prior_run_ids", []), previous.id],
                "diagnostic_reason": diagnostic_reason,
            }
            save_report(manifest_path, manifest)
        if not manifest:
            ledger = read(DATA / "private/pilot-ledger.json", [])
            if len(ledger) + 1 > 24:
                raise NotReady("F1 使用一项诊断预留额度，pilot 与 F1 合计不能超过 24 次。")
            identifier = uid()
            root = DATA / "runs" / identifier / "workspace"
            root.mkdir(parents=True)
            with (root / "sales-simulated.csv").open("w", newline="") as file:
                writer = csv.writer(file)
                writer.writerow(["月份", "产品", "数量", "单价（元）", "退款（元）"])
                writer.writerows(source_rows())
            # Runtime documentation is guidance, not a completed solution or a reference answer.
            roots = Path.home() / ".codex/plugins/cache/openai-primary-runtime"
            docs = [
                roots / "spreadsheets/26.905.11957/skills/spreadsheets/artifact_tool_docs/API_QUICK_START.md",
                roots
                / "presentations/26.905.11957/skills/presentations/artifact_tool_docs/API_QUICK_START.md",
            ]
            for index, document in enumerate(docs):
                if document.exists():
                    shutil.copy2(document, root / f"library-guide-{index + 1}.md")
            (root / "node_modules").symlink_to(RUNTIME / "node/node_modules", target_is_directory=True)
            instruction = "读取 sales-simulated.csv，计算月度和产品销售额、退款和净销售额，制作可编辑销售分析.xlsx 与恰好六页的销售汇报-六页初版.pptx。Excel 使用销售分析和原始销售数据两个工作表；原始销售数据 A2:E19 保留输入，F/G列用公式计算销售额/净销售额，销售分析!D12 用公式汇总总净销售额，加入月度原生图表；销售材料包括可编辑月度和产品原生图表。使用至少两个阶段，完成后会收到三页改版要求。Python可使用已安装openpyxl/python-pptx；若需要公式缓存和重新计算，可用 node_modules 中的 @oai/artifact-tool，参考 library-guide。输出不是静态截图。"
            (root / "public_check.py").write_text(
                "from pathlib import Path\nimport zipfile\nfor name in "
                + repr(NAMES)
                + ":\n assert Path(name).is_file(),name\n with zipfile.ZipFile(name) as package: assert package.testzip() is None\n"
            )
            state = store.create(
                RunState(
                    id=identifier,
                    workspace=str(root),
                    goal=Goal(original=instruction, current=instruction),
                    config=RunConfig(),
                    evaluation_phase="demo",
                    final_checks=[
                        CheckSpec(name="三份可编辑文件完整性", argv=[sys.executable, "public_check.py"])
                    ],
                )
            )
            manifest = {
                "run_id": identifier,
                "data": "synthetic",
                "input_checksum": __import__("hashlib")
                .sha256((root / "sales-simulated.csv").read_bytes())
                .hexdigest(),
                "uses_diagnostic_reserve": True,
            }
            save_report(manifest_path, manifest)
        state = store.get(manifest["run_id"])
        engine = Engine(store, provider)

        async def hook(identifier):
            if (
                not store.has_event(identifier, "f1_revision_requested")
                and inspect_outputs(store.get(identifier).root, revision=False)["passed"]
            ):
                engine.update_goal(identifier, REVISION)
                store.event(identifier, "f1_revision_requested", {"user_feedback": REVISION})

        engine.hook = hook
        if state.phase not in {"completed", "failed", "budget_exhausted", "cancelled"}:
            state = await engine.run(state.id)
        if state.phase == "resource_interrupted":
            raise NotReady("F1 资源中断已保存；恢复额度后再次运行 office-pilot 会继续同一次尝试。")
        destination = DATA / "artifacts" / state.id
        if not destination.exists():
            snapshot(state.root, destination)
        started = time.monotonic()
        verdict = inspect_outputs(destination)
        original_destination = destination
        finalization = None
        if verdict["passed"] and state.phase == "completed":
            receipt = PROJECT / f"outputs/f1-finalization-{state.id}.json"
            finalization = read(receipt)
            if not finalization:
                delivery = PROJECT / "outputs/f1" / ("real-" + state.id)
                finalization = finalize_office_delivery(destination, delivery, NAMES)
                finalization["snapshot"] = digest(delivery)
                save_report(receipt, finalization)
            destination = Path(finalization["output"])
            if digest(destination) != finalization["snapshot"]:
                raise NotReady("办公交付副本已偏离保存的转换记录。")
            verdict = inspect_outputs(destination)
        input_edit = editable_input_check(destination) if verdict["passed"] else {"passed": False}
        real = state.evidence_mode == "real" and store.has_event(state.id, "model_started")
        result = {
            "passed": verdict["passed"]
            and input_edit["passed"]
            and state.phase == "completed"
            and store.has_event(state.id, "f1_revision_requested"),
            "real_agent_run": real,
            "run_id": state.id,
            "output": str(destination),
            "snapshot": digest(destination),
            "original_output": str(original_destination),
            "original_snapshot": digest(original_destination),
            "delivery_finalization": finalization,
            "verdict": verdict,
            "input_edit": input_edit,
            "tokens": state.tokens_used,
            "usage_complete": state.usage_complete,
            "execution_seconds": state.active_seconds,
            "grading_seconds": time.monotonic() - started,
            "visual_review": "pending",
            "target_application_edit_test": "not_run",
            "data": "synthetic",
            "included_in_formal_score": False,
            "prior_attempts": [
                {
                    "run_id": i,
                    "phase": store.get(i).phase,
                    "tokens": store.get(i).tokens_used,
                    "execution_seconds": store.get(i).active_seconds,
                }
                for i in manifest.get("prior_run_ids", [])
            ],
            "all_attempts_tokens": state.tokens_used
            + sum(store.get(i).tokens_used for i in manifest.get("prior_run_ids", [])),
            "diagnostic_reason": manifest.get("diagnostic_reason"),
        }
        save_report(PROJECT / "outputs/f1-validation.json", result)
        return result
    finally:
        await provider.close()
        await broker.close()
        store.close()
